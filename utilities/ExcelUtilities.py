import openpyxl


class ExcelUtilites:

    @staticmethod
    def get_row(file, sheetname):
        wb = openpyxl.load_workbook(file)
        sheet = wb[sheetname]
        return sheet.max_row

    @staticmethod
    def get_column(file, sheetname):
        wb = openpyxl.load_workbook(file)
        sheet = wb[sheetname]
        return sheet.max_column

    @staticmethod
    def read_data(file, sheetname, n_row, n_column):
        wb = openpyxl.load_workbook(file)
        sheet = wb[sheetname]
        return sheet.cell(n_row, n_column).value

    @staticmethod
    def write_data(file, sheetname, n_row, n_column, value):
        wb = openpyxl.load_workbook(file)
        sheet = wb[sheetname]
        sheet.cell(n_row, n_column).value = value
        wb.save(file)